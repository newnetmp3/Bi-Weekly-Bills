from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import load_workbook

from biweekly_bills.database import Database
from biweekly_bills.reports import (
    OTHER_FINANCIAL_REPORTS,
    QUICK_FINANCIAL_REPORTS,
    build_report_bundle,
    export_financial_report,
    export_reports,
)


class ReportExportTests(unittest.TestCase):
    def _fixture(self):
        temp = TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        root = Path(temp.name)
        db = Database(root / "bills.sqlite3")
        db.initialize()

        bill_id = db.upsert_bill(
            name="Verizon",
            cycle="15th",
            default_method="Autopay",
            payment_account="Bills Checking",
            funding_account="Everyday Checking",
            transfer_required=True,
        )
        instance_id = db.upsert_bill_instance(
            year=2026,
            month=9,
            cycle="15th",
            bill_name="Verizon",
            bill_id=bill_id,
            when_label="15th",
            due_cents=21307,
            paid_cents=20000,
            method="Autopay",
            status="Partial",
            source="test",
        )

        db.upsert_bank_account(
            environment="sandbox",
            plaid_account_id="acct-checking",
            name="Checking",
            mask="0157",
            account_type="depository",
            account_subtype="checking",
            current_balance_cents=50000,
            available_balance_cents=45000,
            last_synced_at="2026-09-16T00:00:00+00:00",
        )
        db.set_bills_checking("sandbox", "acct-checking")

        db.upsert_bank_transaction(
            environment="sandbox",
            plaid_transaction_id="tx-verizon",
            plaid_account_id="acct-checking",
            posted_date="2026-09-15",
            authorized_date="2026-09-14",
            merchant_name="Verizon",
            name="VERIZON",
            amount_cents=21307,
            pending=False,
            raw_json="{}",
            last_seen_at="2026-09-16T00:00:00+00:00",
        )
        db.reconcile_transaction(
            environment="sandbox",
            plaid_transaction_id="tx-verizon",
            bill_instance_id=instance_id,
        )
        return root, db

    def test_bundle_contains_month_bills_and_reconciliation_summary(self):
        _, db = self._fixture()
        with patch(
            "biweekly_bills.reports.load_settings",
            return_value=SimpleNamespace(environment="sandbox"),
        ):
            bundle = build_report_bundle(db, 2026, 9)

        self.assertEqual(bundle.month_label, "September 2026")
        self.assertEqual(bundle.bank_label, "Bank data available")
        self.assertEqual(len(bundle.bills), 1)
        self.assertEqual(bundle.bills[0].name, "Verizon")
        self.assertEqual(len(bundle.transactions), 1)
        self.assertEqual(bundle.matched_count, 1)
        self.assertEqual(bundle.unresolved_count, 0)

    def test_bundle_queries_only_selected_transaction_month(self):
        _, db = self._fixture()
        with patch(
            "biweekly_bills.reports.load_settings",
            return_value=SimpleNamespace(environment="sandbox"),
        ), patch.object(
            db,
            "list_bank_transactions",
            wraps=db.list_bank_transactions,
        ) as bank_rows:
            build_report_bundle(db, 2026, 9)

        matching_calls = [
            call
            for call in bank_rows.call_args_list
            if call.kwargs.get("year") == 2026
            and call.kwargs.get("month") == 9
        ]
        self.assertTrue(matching_calls)

    def test_export_all_formats_produces_readable_files(self):
        root, db = self._fixture()
        output = root / "reports"

        with patch(
            "biweekly_bills.reports.load_settings",
            return_value=SimpleNamespace(environment="sandbox"),
        ):
            result = export_reports(
                db,
                2026,
                9,
                formats=("xlsx", "pdf", "ods"),
                output_dir=output,
            )

        self.assertEqual(len(result.paths), 3)
        paths = {path.suffix: path for path in result.paths}
        self.assertEqual(set(paths), {".xlsx", ".pdf", ".ods"})

        for path in result.paths:
            self.assertTrue(path.exists())
            self.assertGreater(path.stat().st_size, 100)

        wb = load_workbook(paths[".xlsx"], data_only=False)
        self.assertEqual(
            wb.sheetnames,
            ["Summary", "Bills", "Funding", "Transactions", "Accounts"],
        )
        self.assertEqual(wb["Summary"]["B1"].value, "September 2026")
        self.assertEqual(wb["Summary"]["B3"].value, "Bank data available")
        self.assertEqual(wb["Bills"]["B2"].value, "Verizon")
        self.assertEqual(wb["Transactions"]["G2"].value, "Matched")

        with paths[".pdf"].open("rb") as handle:
            self.assertEqual(handle.read(5), b"%PDF-")

        self.assertTrue(zipfile.is_zipfile(paths[".ods"]))
        with zipfile.ZipFile(paths[".ods"]) as archive:
            names = set(archive.namelist())
            self.assertIn("mimetype", names)
            self.assertIn("content.xml", names)
            content_xml = archive.read("content.xml").decode("utf-8")
            self.assertIn("September 2026", content_xml)
            self.assertIn("Verizon", content_xml)

    def test_production_bank_data_is_preferred_when_present(self):
        _, db = self._fixture()

        db.upsert_bank_account(
            environment="production",
            plaid_account_id="prod-checking",
            name="NFCU Checking",
            mask="9999",
            account_type="depository",
            account_subtype="checking",
            current_balance_cents=100000,
            available_balance_cents=90000,
            last_synced_at="2026-09-16T00:00:00+00:00",
        )

        funding_bill = db.upsert_bill(
            name="Cox",
            cycle="15th",
            default_method="Autopay",
            payment_account="Bills Checking",
            funding_account="Everyday Checking",
            transfer_required=True,
        )
        db.upsert_bill_instance(
            year=2026,
            month=9,
            cycle="15th",
            bill_name="Cox",
            bill_id=funding_bill,
            due_cents=12060,
            paid_cents=None,
            method="Autopay",
            status="Due",
            source="test",
        )

        with patch(
            "biweekly_bills.reports.load_settings",
            return_value=SimpleNamespace(environment="production"),
        ):
            bundle = build_report_bundle(db, 2026, 9)
        self.assertEqual(bundle.bank_environment, "production")
        self.assertEqual(bundle.bank_label, "Bank data available")
        self.assertEqual(len(bundle.accounts), 1)
        self.assertEqual(bundle.accounts[0].name, "NFCU Checking")
        self.assertEqual(len(bundle.transactions), 0)
        self.assertEqual(len(bundle.funding_items), 1)
        self.assertEqual(bundle.funding_items[0].bill_name, "Cox")
        self.assertEqual(bundle.funding_items[0].remaining_cents, 12060)

    def test_financial_report_catalog_exports_readable_workbooks(self):
        root, db = self._fixture()
        output = root / "financial-reports"

        report_keys = [
            key
            for key, _label in (
                QUICK_FINANCIAL_REPORTS + OTHER_FINANCIAL_REPORTS
            )
        ]
        self.assertEqual(len(report_keys), 8)

        with patch(
            "biweekly_bills.reports.load_settings",
            return_value=SimpleNamespace(environment="sandbox"),
        ):
            for report_key in report_keys:
                with self.subTest(report_key=report_key):
                    result = export_financial_report(
                        db,
                        2026,
                        9,
                        report_key,
                        output_dir=output,
                    )
                    self.assertEqual(len(result.paths), 1)
                    path = result.paths[0]
                    self.assertEqual(path.suffix, ".xlsx")
                    self.assertTrue(path.exists())
                    self.assertGreater(path.stat().st_size, 100)

                    workbook = load_workbook(path, data_only=False)
                    self.assertTrue(workbook.sheetnames)
                    self.assertTrue(
                        any(
                            cell.value is not None
                            for row in workbook.active.iter_rows()
                            for cell in row
                        )
                    )

    def test_production_mode_never_falls_back_to_sandbox_bank_data(self):
        _, db = self._fixture()

        with patch(
            "biweekly_bills.reports.load_settings",
            return_value=SimpleNamespace(environment="production"),
        ):
            bundle = build_report_bundle(db, 2026, 9)

        self.assertIsNone(bundle.bank_environment)
        self.assertEqual(bundle.bank_label, "No bank data")
        self.assertEqual(bundle.accounts, ())
        self.assertEqual(bundle.transactions, ())
        self.assertEqual(bundle.matched_count, 0)
        self.assertEqual(bundle.ignored_count, 0)
        self.assertEqual(bundle.unresolved_count, 0)



if __name__ == "__main__":
    unittest.main()

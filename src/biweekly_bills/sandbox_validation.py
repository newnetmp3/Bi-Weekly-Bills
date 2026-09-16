from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import sqlite3
import tempfile
from typing import Iterable
import zipfile

from openpyxl import load_workbook

from .backups import BackupManager, validate_sqlite
from .secure_store import APP_DIR
from .bank_sync import sandbox_connection_status, sync_sandbox_to_sqlite
from .database import Database, utc_now
from .funding import build_funding_plan
from .reconciliation import accept_match, undo_reconciliation
from .reports import export_reports


SANDBOX_VALIDATION_MARKER = APP_DIR / "sandbox_validation_pass.json"

VALIDATION_TABLES = (
    "metadata",
    "bills",
    "pay_periods",
    "bill_instances",
    "bank_accounts",
    "bank_transactions",
    "sync_state",
    "transaction_reconciliations",
    "workbook_imports",
)


@dataclass(frozen=True)
class ValidationCheck:
    name: str
    status: str
    detail: str


@dataclass(frozen=True)
class SandboxValidationReport:
    started_at: datetime
    finished_at: datetime
    checks: tuple[ValidationCheck, ...]

    @property
    def passed(self) -> bool:
        return not any(check.status == "FAIL" for check in self.checks)

    @property
    def pass_count(self) -> int:
        return sum(check.status == "PASS" for check in self.checks)

    @property
    def warn_count(self) -> int:
        return sum(check.status == "WARN" for check in self.checks)

    @property
    def fail_count(self) -> int:
        return sum(check.status == "FAIL" for check in self.checks)

    @property
    def skip_count(self) -> int:
        return sum(check.status == "SKIP" for check in self.checks)


def _logical_fingerprint(database: Database) -> str:
    digest = hashlib.sha256()
    uri = f"file:{database.path.resolve()}?mode=ro"
    conn = sqlite3.connect(uri, uri=True)
    try:
        for table in VALIDATION_TABLES:
            digest.update(table.encode("utf-8"))
            try:
                rows = conn.execute(
                    f"SELECT * FROM {table} ORDER BY rowid"
                ).fetchall()
            except sqlite3.OperationalError:
                rows = []
            for row in rows:
                digest.update(repr(tuple(row)).encode("utf-8"))
    finally:
        conn.close()
    return digest.hexdigest()


def _clone_database(source: Database, target_path: Path) -> Database:
    target_path.parent.mkdir(parents=True, exist_ok=True)
    source_conn = sqlite3.connect(source.path)
    target_conn = sqlite3.connect(target_path)
    try:
        source_conn.backup(target_conn)
    finally:
        target_conn.close()
        source_conn.close()

    clone = Database(target_path)
    clone.initialize()
    return clone


def _remove_production_data(clone: Database) -> None:
    with clone.transaction() as conn:
        conn.execute(
            "DELETE FROM transaction_reconciliations WHERE environment='production'"
        )
        conn.execute("DELETE FROM bank_transactions WHERE environment='production'")
        conn.execute("DELETE FROM bank_accounts WHERE environment='production'")
        conn.execute("DELETE FROM sync_state WHERE environment='production'")


def _check_report_files(paths: Iterable[Path]) -> str:
    by_suffix = {path.suffix.casefold(): path for path in paths}
    expected = {".xlsx", ".pdf", ".ods"}
    if set(by_suffix) != expected:
        raise RuntimeError(
            f"Expected XLSX/PDF/ODS outputs, got {sorted(by_suffix)}"
        )

    workbook = load_workbook(by_suffix[".xlsx"], read_only=True, data_only=False)
    try:
        expected_sheets = [
            "Summary",
            "Bills",
            "Funding",
            "Transactions",
            "Accounts",
        ]
        if workbook.sheetnames != expected_sheets:
            raise RuntimeError(
                f"Excel sheets do not match expected layout: {workbook.sheetnames}"
            )
    finally:
        workbook.close()

    with by_suffix[".pdf"].open("rb") as handle:
        if handle.read(5) != b"%PDF-":
            raise RuntimeError("PDF output does not have a valid PDF signature.")

    if not zipfile.is_zipfile(by_suffix[".ods"]):
        raise RuntimeError("ODS output is not a valid OpenDocument package.")
    with zipfile.ZipFile(by_suffix[".ods"]) as archive:
        names = set(archive.namelist())
        if "mimetype" not in names or "content.xml" not in names:
            raise RuntimeError("ODS output is missing required package members.")

    return ", ".join(path.name for path in paths)


def load_sandbox_validation_marker() -> dict:
    if not SANDBOX_VALIDATION_MARKER.exists():
        return {}
    try:
        with SANDBOX_VALIDATION_MARKER.open("r", encoding="utf-8") as handle:
            data = json.load(handle)
    except (OSError, json.JSONDecodeError):
        return {}
    return data if isinstance(data, dict) else {}


def record_sandbox_validation_pass(report: SandboxValidationReport) -> Path:
    if not report.passed or report.fail_count:
        raise ValueError("Only a successful Sandbox validation can be recorded.")

    APP_DIR.mkdir(parents=True, exist_ok=True)
    os.chmod(APP_DIR, 0o700)
    payload = {
        "passed": True,
        "started_at": report.started_at.isoformat(),
        "finished_at": report.finished_at.isoformat(),
        "pass_count": report.pass_count,
        "warn_count": report.warn_count,
        "skip_count": report.skip_count,
    }
    temp = SANDBOX_VALIDATION_MARKER.with_suffix(".json.tmp")
    with temp.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2)
        handle.write("\n")
        handle.flush()
        os.fsync(handle.fileno())
    os.chmod(temp, 0o600)
    os.replace(temp, SANDBOX_VALIDATION_MARKER)
    os.chmod(SANDBOX_VALIDATION_MARKER, 0o600)
    return SANDBOX_VALIDATION_MARKER


def run_sandbox_validation(
    database: Database,
    *,
    perform_live_sync: bool = True,
) -> SandboxValidationReport:
    """Run the full Sandbox QA workflow without mutating the working database.

    The working SQLite database is cloned first. Any live Plaid Sandbox sync,
    reconciliation round-trip, funding simulation, backup/restore mutation, and
    report generation happens only inside the temporary clone. Production
    credentials/APIs are never requested by this routine.
    """

    started = datetime.now(timezone.utc)
    checks: list[ValidationCheck] = []

    def add(name: str, status: str, detail: str) -> None:
        checks.append(ValidationCheck(name=name, status=status, detail=detail))

    before_fingerprint = _logical_fingerprint(database)

    production_accounts = database.list_bank_accounts("production")
    if production_accounts:
        add(
            "Production isolation",
            "WARN",
            f"{len(production_accounts)} Production account row(s) exist locally. "
            "They will be removed from the temporary validation clone; Production APIs are not contacted.",
        )
    else:
        add(
            "Production isolation",
            "PASS",
            "No Production bank accounts are present. Validation is Sandbox-only.",
        )

    connection = sandbox_connection_status()
    sandbox_ready = (
        str(connection.get("configured_environment") or "").casefold() == "sandbox"
        and bool(connection.get("keys_configured"))
        and bool(connection.get("credential_present"))
    )
    if perform_live_sync:
        if sandbox_ready:
            item_id = str(connection.get("item_id") or "")
            detail = "Sandbox keys and persisted access token are available."
            if item_id:
                detail += f" Item {item_id[:12]}…"
            add("Sandbox connection", "PASS", detail)
        else:
            missing: list[str] = []
            if str(connection.get("configured_environment") or "").casefold() != "sandbox":
                missing.append("PLAID_ENV is not sandbox")
            if not connection.get("keys_configured"):
                missing.append("API keys are not configured")
            if not connection.get("credential_present"):
                missing.append("Sandbox access token is missing")
            add(
                "Sandbox connection",
                "FAIL",
                "; ".join(missing) or "Sandbox connection is not ready.",
            )
    else:
        add(
            "Sandbox connection",
            "SKIP",
            "Live Plaid connection check disabled for deterministic/offline validation.",
        )

    with tempfile.TemporaryDirectory(prefix="biweekly-bills-validation-") as temp_name:
        temp_root = Path(temp_name)
        clone = _clone_database(database, temp_root / "validation.sqlite3")
        _remove_production_data(clone)

        if validate_sqlite(clone.path):
            add(
                "Isolated database clone",
                "PASS",
                "Working SQLite data was cloned and passed integrity_check.",
            )
        else:
            add(
                "Isolated database clone",
                "FAIL",
                "Temporary validation clone failed SQLite integrity verification.",
            )

        if perform_live_sync and sandbox_ready:
            try:
                sync = sync_sandbox_to_sqlite(
                    clone,
                    persist_local_cursor=False,
                )
                add(
                    "Live Sandbox sync",
                    "PASS",
                    f"{sync.account_count} account(s), {sync.transaction_count} stored transaction(s), "
                    f"+{sync.added_count} / ~{sync.modified_count} / -{sync.removed_count}; "
                    f"{sync.update_status}. Working cursor was not advanced.",
                )
            except Exception as exc:
                add("Live Sandbox sync", "FAIL", str(exc))
        elif perform_live_sync:
            add(
                "Live Sandbox sync",
                "SKIP",
                "Skipped because the Sandbox connection prerequisite failed.",
            )
        else:
            add(
                "Live Sandbox sync",
                "SKIP",
                "Live Plaid API call disabled for this validation run.",
            )

        sandbox_accounts = clone.list_bank_accounts("sandbox")
        sandbox_transactions = clone.list_bank_transactions("sandbox", limit=100000)
        if sandbox_accounts and sandbox_transactions:
            add(
                "Sandbox cache",
                "PASS",
                f"{len(sandbox_accounts)} account(s) and {len(sandbox_transactions)} transaction(s) "
                "are available in the isolated SQLite clone.",
            )
        else:
            add(
                "Sandbox cache",
                "FAIL",
                f"Validation requires cached Sandbox accounts and transactions; found "
                f"{len(sandbox_accounts)} account(s) and {len(sandbox_transactions)} transaction(s).",
            )

        bills_account = next(
            (row for row in sandbox_accounts if int(row["is_bills_checking"] or 0)),
            None,
        )
        if bills_account is None:
            add(
                "Sandbox Bills account role",
                "WARN",
                "No Sandbox account is designated Bills Checking. Other tests continue.",
            )
        else:
            label = str(bills_account["name"] or "(unnamed)")
            if bills_account["mask"]:
                label += f" ••••{bills_account['mask']}"
            add(
                "Sandbox Bills account role",
                "PASS",
                f"{label} is designated Bills Checking in the validation clone.",
            )

        today = date.today()
        validation_date = today.replace(day=5)
        account_id = (
            str(sandbox_accounts[0]["plaid_account_id"])
            if sandbox_accounts
            else "validation-sandbox-account"
        )
        if not sandbox_accounts:
            clone.upsert_bank_account(
                environment="sandbox",
                plaid_account_id=account_id,
                name="Validation Checking",
                mask="0001",
                account_type="depository",
                account_subtype="checking",
                current_balance_cents=50000,
                available_balance_cents=50000,
                last_synced_at=utc_now(),
            )

        try:
            bill_id = clone.upsert_bill(
                name="__Validation Reconciliation Bill__",
                cycle="1st",
                default_method="Validation",
            )
            instance_id = clone.upsert_bill_instance(
                year=validation_date.year,
                month=validation_date.month,
                cycle="1st",
                bill_name="__Validation Reconciliation Bill__",
                bill_id=bill_id,
                due_cents=12345,
                paid_cents=None,
                status="Due",
                source="validation",
            )
            clone.upsert_bank_transaction(
                environment="sandbox",
                plaid_transaction_id="__validation_reconciliation_tx__",
                plaid_account_id=account_id,
                posted_date=validation_date.isoformat(),
                authorized_date=validation_date.isoformat(),
                merchant_name="Validation Merchant",
                name="Validation Payment",
                amount_cents=12345,
                pending=False,
                raw_json="{}",
                last_seen_at=utc_now(),
            )
            tx = clone.get_bank_transaction(
                "sandbox",
                "__validation_reconciliation_tx__",
            )
            if tx is None:
                raise RuntimeError("Validation transaction was not created.")

            accept_match(clone, tx, instance_id)
            matched = clone.get_bill_instance(instance_id)
            if (
                matched is None
                or int(matched["paid_cents"] or 0) != 12345
                or str(matched["status"] or "") != "Paid"
                or str(matched["source"] or "") != "bank-reconciled"
            ):
                raise RuntimeError("Accept Match did not update Paid/Status/source correctly.")

            tx = clone.get_bank_transaction(
                "sandbox",
                "__validation_reconciliation_tx__",
            )
            if tx is None:
                raise RuntimeError("Matched validation transaction disappeared.")
            undo_reconciliation(clone, tx)
            restored = clone.get_bill_instance(instance_id)
            if (
                restored is None
                or restored["paid_cents"] is not None
                or str(restored["status"] or "") != "Due"
                or str(restored["source"] or "") != "validation"
            ):
                raise RuntimeError("Undo did not restore the pre-match bill state.")

            add(
                "Reconciliation round-trip",
                "PASS",
                "Accept Match updated Paid/Status/source and Undo restored the exact prior state.",
            )
        except Exception as exc:
            add("Reconciliation round-trip", "FAIL", str(exc))

        try:
            gate = build_funding_plan(
                clone,
                validation_date.year,
                validation_date.month,
            )
            if gate.environment is not None:
                raise RuntimeError(
                    "Sandbox data incorrectly activated Production funding calculations."
                )
            add(
                "Funding Sandbox gate",
                "PASS",
                "Sandbox accounts do not activate real Payment/Funding transfer planning.",
            )
        except Exception as exc:
            add("Funding Sandbox gate", "FAIL", str(exc))

        try:
            report_dir = temp_root / "reports"
            exported = export_reports(
                clone,
                validation_date.year,
                validation_date.month,
                formats=("xlsx", "pdf", "ods"),
                output_dir=report_dir,
            )
            detail = _check_report_files(exported.paths)
            add(
                "Report export round-trip",
                "PASS",
                f"Generated and reopened Excel/PDF/ODS successfully: {detail}",
            )
        except Exception as exc:
            add("Report export round-trip", "FAIL", str(exc))

        try:
            clone.upsert_bank_account(
                environment="production",
                plaid_account_id="__validation_production_bills__",
                name="Validation Bills Checking",
                mask="9001",
                account_type="depository",
                account_subtype="checking",
                current_balance_cents=0,
                available_balance_cents=0,
                last_synced_at=utc_now(),
            )
            clone.set_bills_checking(
                "production",
                "__validation_production_bills__",
            )
            funding_bill = clone.upsert_bill(
                name="__Validation Funding Bill__",
                cycle="1st",
                payment_account="Bills Checking",
                funding_account="Validation Everyday Checking",
                transfer_required=True,
            )
            clone.upsert_bill_instance(
                year=2099,
                month=12,
                cycle="1st",
                bill_name="__Validation Funding Bill__",
                bill_id=funding_bill,
                due_cents=43210,
                paid_cents=10000,
                status="Partial",
                source="validation",
            )
            plan = build_funding_plan(
                clone,
                2099,
                12,
                today=date(2099, 11, 1),
            )
            if (
                plan.first_required_cents != 33210
                or plan.total_required_cents != 33210
                or plan.total_transfer_cents != 33210
            ):
                raise RuntimeError(
                    "Funding calculation did not produce the expected $332.10 remaining transfer."
                )
            add(
                "Funding calculation",
                "PASS",
                "Isolated Production-like fixture calculated $332.10 remaining/transfer without touching real Production.",
            )
        except Exception as exc:
            add("Funding calculation", "FAIL", str(exc))

        try:
            clone_backups = BackupManager(
                clone,
                backup_dir=temp_root / "backups",
                retention=5,
            )
            known_good = clone_backups.create_backup("validation-known-good")
            clone.upsert_bill(
                name="__Validation Restore Sentinel__",
                cycle="15th",
            )
            if not any(
                row["name"] == "__Validation Restore Sentinel__"
                for row in clone.list_bills(active_only=False)
            ):
                raise RuntimeError("Restore sentinel was not written before restore.")

            restored, safety = clone_backups.restore(known_good.path)
            if any(
                row["name"] == "__Validation Restore Sentinel__"
                for row in clone.list_bills(active_only=False)
            ):
                raise RuntimeError("Restore did not roll back the sentinel mutation.")
            if not restored.integrity_ok or not safety.integrity_ok:
                raise RuntimeError("Restored or pre-restore safety backup failed integrity verification.")
            add(
                "Backup / restore round-trip",
                "PASS",
                "Verified backup restored the clone and preserved a verified pre-restore safety snapshot.",
            )
        except Exception as exc:
            add("Backup / restore round-trip", "FAIL", str(exc))

    after_fingerprint = _logical_fingerprint(database)
    if after_fingerprint == before_fingerprint:
        add(
            "Working database unchanged",
            "PASS",
            "Full validation completed without changing the working SQLite database.",
        )
    else:
        add(
            "Working database unchanged",
            "FAIL",
            "Working SQLite data changed during isolated validation.",
        )

    finished = datetime.now(timezone.utc)
    return SandboxValidationReport(
        started_at=started,
        finished_at=finished,
        checks=tuple(checks),
    )

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from odf import table as odf_table
from odf import text as odf_text
from odf.opendocument import OpenDocumentSpreadsheet
from odf.style import Style, TableCellProperties, TextProperties
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .database import Database, MonthSummary
from .funding import FundingPlan, build_funding_plan, funding_lines
from .settings import load_settings


MONTH_NAMES = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]


@dataclass(frozen=True)
class ReportBillRow:
    cycle: str
    name: str
    when_label: str | None
    due_cents: int
    paid_cents: int
    remaining_cents: int
    method: str | None
    status: str | None
    payment_account: str | None
    transfer_source: str | None
    transfer_required: str
    extra_short: str | None


@dataclass(frozen=True)
class ReportFundingRow:
    cycle: str
    bill_name: str
    due_cents: int
    paid_cents: int
    remaining_cents: int
    payment_account: str | None
    transfer_source: str | None


@dataclass(frozen=True)
class ReportAccountRow:
    name: str
    mask: str | None
    account_type: str | None
    account_subtype: str | None
    current_balance_cents: int | None
    available_balance_cents: int | None
    is_bills_checking: bool


@dataclass(frozen=True)
class ReportTransactionRow:
    date: str
    merchant: str
    account: str
    amount_cents: int
    pending: bool
    reconciliation: str
    reconciled_bill: str | None


@dataclass(frozen=True)
class ReportBundle:
    year: int
    month: int
    generated_at: datetime
    month_summary: MonthSummary
    first_summary: MonthSummary
    fifteenth_summary: MonthSummary
    funding: FundingPlan
    bank_environment: str | None
    bills: tuple[ReportBillRow, ...]
    funding_items: tuple[ReportFundingRow, ...]
    accounts: tuple[ReportAccountRow, ...]
    transactions: tuple[ReportTransactionRow, ...]
    matched_count: int
    ignored_count: int
    funding_validated_count: int
    unresolved_count: int

    @property
    def month_label(self) -> str:
        return f"{MONTH_NAMES[self.month - 1]} {self.year}"

    @property
    def bank_label(self) -> str:
        if self.bank_environment is not None:
            return "Bank data available"
        return "No bank data"


@dataclass(frozen=True)
class ExportResult:
    paths: tuple[Path, ...]


def default_reports_dir(database: Database) -> Path:
    return database.path.parent / "reports"


def _report_environment(database: Database) -> str | None:
    configured = load_settings(require_keys=False).environment
    if configured == "production":
        return "production" if database.list_bank_accounts("production") else None
    return "sandbox" if database.list_bank_accounts("sandbox") else None


def _joined_bill_account_label(row: Any, prefix: str) -> str | None:
    name = row[f"{prefix}_name"]
    mask = row[f"{prefix}_mask"]
    if not name:
        return None
    label = str(name)
    if mask:
        label += f" ••••{mask}"
    return label


def build_report_bundle(database: Database, year: int, month: int) -> ReportBundle:
    bills: list[ReportBillRow] = []
    for row in database.list_month_instances(year, month):
        due = int(row["due_cents"] or 0)
        paid = int(row["paid_cents"] or 0)
        transfer = row["transfer_required"]
        transfer_label = (
            "Not specified"
            if transfer is None
            else ("Yes" if int(transfer) else "No")
        )
        bills.append(
            ReportBillRow(
                cycle=str(row["cycle"]),
                name=str(row["bill_name_snapshot"]),
                when_label=row["when_label"],
                due_cents=due,
                paid_cents=paid,
                remaining_cents=max(due - paid, 0),
                method=row["method"],
                status=row["status"],
                payment_account=(
                    row["payment_account_snapshot"]
                    or _joined_bill_account_label(row, "payment_account")
                    or row["payment_account"]
                ),
                transfer_source=(
                    _joined_bill_account_label(row, "transfer_source_account")
                    or row["funding_account"]
                ),
                transfer_required=transfer_label,
                extra_short=row["extra_short"],
            )
        )

    environment = _report_environment(database)
    funding_plan = build_funding_plan(database, year, month)
    funding_items: list[ReportFundingRow] = []
    if funding_plan.environment is not None:
        lines, _ = funding_lines(database, year, month)
        for line in lines:
            funding_items.append(
                ReportFundingRow(
                    cycle=line.cycle,
                    bill_name=line.bill_name,
                    due_cents=line.due_cents,
                    paid_cents=line.paid_cents,
                    remaining_cents=line.remaining_cents,
                    payment_account=line.payment_account_label,
                    transfer_source=line.transfer_source_label,
                )
            )

    accounts: list[ReportAccountRow] = []
    transactions: list[ReportTransactionRow] = []
    matched = ignored = funding_validated = unresolved = 0

    if environment is not None:
        for row in database.list_bank_accounts(environment):
            accounts.append(
                ReportAccountRow(
                    name=str(row["name"] or "(unnamed account)"),
                    mask=str(row["mask"] or "") or None,
                    account_type=row["account_type"],
                    account_subtype=row["account_subtype"],
                    current_balance_cents=(
                        None if row["current_balance_cents"] is None
                        else int(row["current_balance_cents"])
                    ),
                    available_balance_cents=(
                        None if row["available_balance_cents"] is None
                        else int(row["available_balance_cents"])
                    ),
                    is_bills_checking=bool(row["is_bills_checking"]),
                )
            )

        for row in database.list_bank_transactions(
            environment,
            limit=10000,
            year=year,
            month=month,
        ):
            raw_date = str(row["posted_date"] or row["authorized_date"] or "")

            disposition = row["reconciliation_disposition"]
            funding_scope = row["funding_validation_scope"]
            if funding_scope:
                scope_label = (
                    "Whole month"
                    if funding_scope == "month"
                    else f"{funding_scope} pay period"
                )
                reconciliation = f"Funding validated — {scope_label}"
                funding_validated += 1
            elif disposition == "matched":
                reconciliation = "Matched"
                matched += 1
            elif disposition == "ignored":
                reconciliation = "Ignored"
                ignored += 1
            else:
                reconciliation = "Unresolved"
                unresolved += 1

            account = str(row["account_name"] or "(unknown account)")
            if row["account_mask"]:
                account += f" ••••{row['account_mask']}"

            transactions.append(
                ReportTransactionRow(
                    date=raw_date,
                    merchant=str(
                        row["merchant_name"] or row["name"] or "(unnamed transaction)"
                    ),
                    account=account,
                    amount_cents=int(row["amount_cents"]),
                    pending=bool(row["pending"]),
                    reconciliation=reconciliation,
                    reconciled_bill=(
                        str(row["reconciled_bill_name"])
                        if row["reconciled_bill_name"]
                        else None
                    ),
                )
            )

    return ReportBundle(
        year=year,
        month=month,
        generated_at=datetime.now(timezone.utc),
        month_summary=database.month_summary(year, month),
        first_summary=database.cycle_summary(year, month, "1st"),
        fifteenth_summary=database.cycle_summary(year, month, "15th"),
        funding=funding_plan,
        bank_environment=environment,
        bills=tuple(bills),
        funding_items=tuple(funding_items),
        accounts=tuple(accounts),
        transactions=tuple(transactions),
        matched_count=matched,
        ignored_count=ignored,
        funding_validated_count=funding_validated,
        unresolved_count=unresolved,
    )


def _stem(bundle: ReportBundle) -> str:
    stamp = bundle.generated_at.strftime("%Y%m%dT%H%M%SZ")
    return f"biweekly-bills-{bundle.year:04d}-{bundle.month:02d}-{stamp}"


def export_reports(
    database: Database,
    year: int,
    month: int,
    *,
    formats: tuple[str, ...] = ("xlsx", "pdf", "ods"),
    output_dir: Path | str | None = None,
) -> ExportResult:
    bundle = build_report_bundle(database, year, month)
    destination = Path(output_dir) if output_dir is not None else default_reports_dir(database)
    destination.mkdir(parents=True, exist_ok=True)
    stem = _stem(bundle)

    paths: list[Path] = []
    for fmt in formats:
        fmt = fmt.casefold().lstrip(".")
        if fmt == "xlsx":
            path = destination / f"{stem}.xlsx"
            export_excel(bundle, path)
        elif fmt == "pdf":
            path = destination / f"{stem}.pdf"
            export_pdf(bundle, path)
        elif fmt == "ods":
            path = destination / f"{stem}.ods"
            export_ods(bundle, path)
        else:
            raise ValueError(f"Unsupported report format: {fmt}")
        paths.append(path)

    return ExportResult(paths=tuple(paths))


def _money(cents: int | None) -> float | None:
    return None if cents is None else cents / 100.0


def _autowidth(ws) -> None:
    for column_cells in ws.columns:
        length = max(
            (len("" if cell.value is None else str(cell.value)) for cell in column_cells),
            default=0,
        )
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            max(length + 2, 10), 42
        )


def _xlsx_header(ws, row: int, columns: int) -> None:
    fill = PatternFill("solid", fgColor="10182B")
    font = Font(color="C8FF3D", bold=True)
    for col in range(1, columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center")


def export_excel(bundle: ReportBundle, path: Path) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    summary_rows = [
        ("Bi-Weekly Bills Report", bundle.month_label),
        ("Generated", bundle.generated_at.astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")),
        ("Bank data", bundle.bank_label),
        ("", ""),
        ("Month scheduled", _money(bundle.month_summary.due_cents)),
        ("Month paid", _money(bundle.month_summary.paid_cents)),
        ("Month remaining", _money(bundle.month_summary.remaining_cents)),
        ("1st scheduled", _money(bundle.first_summary.due_cents)),
        ("1st paid", _money(bundle.first_summary.paid_cents)),
        ("15th scheduled", _money(bundle.fifteenth_summary.due_cents)),
        ("15th paid", _money(bundle.fifteenth_summary.paid_cents)),
        ("", ""),
        ("Bills funding gross remaining", _money(bundle.funding.total_required_cents)),
        ("Transfer for 1st", _money(bundle.funding.first_transfer_cents)),
        ("Transfer for 15th", _money(bundle.funding.fifteenth_transfer_cents)),
        ("Month transfer", _money(bundle.funding.total_transfer_cents)),
        ("Bills balance", _money(bundle.funding.available_balance_cents)),
        ("Funding items needing review", bundle.funding.review_bill_count),
        ("", ""),
        ("Transactions", len(bundle.transactions)),
        ("Matched", bundle.matched_count),
        ("Ignored", bundle.ignored_count),
        ("Funding transfers validated", bundle.funding_validated_count),
        ("Unresolved", bundle.unresolved_count),
    ]
    for row in summary_rows:
        summary.append(row)
    summary["A1"].font = Font(bold=True, size=16, color="C8FF3D")
    summary["A1"].fill = PatternFill("solid", fgColor="10182B")
    summary["B1"].font = Font(bold=True, size=14, color="FFFFFF")
    summary["B1"].fill = PatternFill("solid", fgColor="10182B")
    for row in range(5, 18):
        summary.cell(row=row, column=2).number_format = '$#,##0.00;[Red]-$#,##0.00'
    _autowidth(summary)

    bills = wb.create_sheet("Bills")
    bill_headers = [
        "Cycle", "Bill", "When", "Due", "Paid", "Remaining", "Method", "Status",
        "Payment Account", "Transfer Source", "Transfer Required", "Extra / Short",
    ]
    bills.append(bill_headers)
    _xlsx_header(bills, 1, len(bill_headers))
    for row in bundle.bills:
        bills.append([
            row.cycle, row.name, row.when_label, _money(row.due_cents),
            _money(row.paid_cents), _money(row.remaining_cents), row.method,
            row.status, row.payment_account, row.transfer_source,
            row.transfer_required, row.extra_short,
        ])
    for row in bills.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = '$#,##0.00;[Red]-$#,##0.00'
    bills.freeze_panes = "A2"
    bills.auto_filter.ref = bills.dimensions
    _autowidth(bills)

    funding = wb.create_sheet("Funding")
    funding_headers = [
        "Cycle", "Bill", "Due", "Paid", "Remaining", "Payment Account", "Transfer Source",
    ]
    funding.append(funding_headers)
    _xlsx_header(funding, 1, len(funding_headers))
    for row in bundle.funding_items:
        funding.append([
            row.cycle,
            row.bill_name,
            _money(row.due_cents),
            _money(row.paid_cents),
            _money(row.remaining_cents),
            row.payment_account,
            row.transfer_source,
        ])
    for row in funding.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = '$#,##0.00;[Red]-$#,##0.00'
    funding.freeze_panes = "A2"
    funding.auto_filter.ref = funding.dimensions
    _autowidth(funding)

    tx = wb.create_sheet("Transactions")
    tx_headers = [
        "Date", "Merchant / Description", "Account", "Amount", "Flow",
        "State", "Reconciliation", "Matched Bill",
    ]
    tx.append(tx_headers)
    _xlsx_header(tx, 1, len(tx_headers))
    for row in bundle.transactions:
        tx.append([
            row.date, row.merchant, row.account, _money(row.amount_cents),
            "Outflow" if row.amount_cents >= 0 else "Inflow",
            "Pending" if row.pending else "Posted",
            row.reconciliation, row.reconciled_bill,
        ])
    for cell in tx["D"][1:]:
        cell.number_format = '$#,##0.00;[Red]-$#,##0.00'
    tx.freeze_panes = "A2"
    tx.auto_filter.ref = tx.dimensions
    _autowidth(tx)

    accounts = wb.create_sheet("Accounts")
    account_headers = [
        "Account", "Mask", "Type", "Subtype", "Current", "Available", "Bills Checking",
    ]
    accounts.append(account_headers)
    _xlsx_header(accounts, 1, len(account_headers))
    for row in bundle.accounts:
        accounts.append([
            row.name,
            f"••••{row.mask}" if row.mask else None,
            row.account_type,
            row.account_subtype,
            _money(row.current_balance_cents),
            _money(row.available_balance_cents),
            "Yes" if row.is_bills_checking else "No",
        ])
    for col in ("E", "F"):
        for cell in accounts[col][1:]:
            cell.number_format = '$#,##0.00;[Red]-$#,##0.00'
    accounts.freeze_panes = "A2"
    accounts.auto_filter.ref = accounts.dimensions
    _autowidth(accounts)

    wb.save(path)


def _ods_text_cell(value: Any, *, style_name: str | None = None):
    kwargs = {}
    if style_name:
        kwargs["stylename"] = style_name
    cell = odf_table.TableCell(**kwargs)
    cell.addElement(odf_text.P(text="" if value is None else str(value)))
    return cell


def _ods_currency_cell(cents: int | None, *, style_name: str | None = None):
    if cents is None:
        return _ods_text_cell("", style_name=style_name)
    kwargs = {}
    if style_name:
        kwargs["stylename"] = style_name
    value = cents / 100.0
    cell = odf_table.TableCell(
        valuetype="currency",
        currency="USD",
        value=value,
        **kwargs,
    )
    cell.addElement(odf_text.P(text="$" + f"{value:,.2f}"))
    return cell


def _ods_add_row(sheet, values, *, header_style=None):
    row = odf_table.TableRow()
    for value in values:
        if isinstance(value, tuple) and value and value[0] == "currency":
            row.addElement(_ods_currency_cell(value[1], style_name=header_style))
        else:
            row.addElement(_ods_text_cell(value, style_name=header_style))
    sheet.addElement(row)


def export_ods(bundle: ReportBundle, path: Path) -> None:
    doc = OpenDocumentSpreadsheet()
    header_style = Style(name="HeaderCell", family="table-cell")
    header_style.addElement(
        TableCellProperties(backgroundcolor="#10182B", padding="0.05in")
    )
    header_style.addElement(TextProperties(color="#C8FF3D", fontweight="bold"))
    doc.styles.addElement(header_style)

    summary = odf_table.Table(name="Summary")
    _ods_add_row(
        summary,
        ["Bi-Weekly Bills Report", bundle.month_label],
        header_style="HeaderCell",
    )
    for label, value in [
        ("Generated", bundle.generated_at.astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")),
        ("Bank data", bundle.bank_label),
        ("Month scheduled", ("currency", bundle.month_summary.due_cents)),
        ("Month paid", ("currency", bundle.month_summary.paid_cents)),
        ("Month remaining", ("currency", bundle.month_summary.remaining_cents)),
        ("1st scheduled", ("currency", bundle.first_summary.due_cents)),
        ("1st paid", ("currency", bundle.first_summary.paid_cents)),
        ("15th scheduled", ("currency", bundle.fifteenth_summary.due_cents)),
        ("15th paid", ("currency", bundle.fifteenth_summary.paid_cents)),
        ("Bills funding gross remaining", ("currency", bundle.funding.total_required_cents)),
        ("Transfer for 1st", ("currency", bundle.funding.first_transfer_cents)),
        ("Transfer for 15th", ("currency", bundle.funding.fifteenth_transfer_cents)),
        ("Month transfer", ("currency", bundle.funding.total_transfer_cents)),
        ("Bills balance", ("currency", bundle.funding.available_balance_cents)),
        ("Transactions", len(bundle.transactions)),
        ("Matched", bundle.matched_count),
        ("Ignored", bundle.ignored_count),
        ("Funding transfers validated", bundle.funding_validated_count),
        ("Unresolved", bundle.unresolved_count),
    ]:
        _ods_add_row(summary, [label, value])
    doc.spreadsheet.addElement(summary)

    bills = odf_table.Table(name="Bills")
    bill_headers = [
        "Cycle", "Bill", "When", "Due", "Paid", "Remaining", "Method", "Status",
        "Payment Account", "Transfer Source", "Transfer Required", "Extra / Short",
    ]
    _ods_add_row(bills, bill_headers, header_style="HeaderCell")
    for row in bundle.bills:
        _ods_add_row(bills, [
            row.cycle, row.name, row.when_label,
            ("currency", row.due_cents), ("currency", row.paid_cents),
            ("currency", row.remaining_cents), row.method, row.status,
            row.payment_account, row.transfer_source, row.transfer_required,
            row.extra_short,
        ])
    doc.spreadsheet.addElement(bills)

    funding = odf_table.Table(name="Funding")
    funding_headers = [
        "Cycle", "Bill", "Due", "Paid", "Remaining", "Payment Account", "Transfer Source",
    ]
    _ods_add_row(funding, funding_headers, header_style="HeaderCell")
    for row in bundle.funding_items:
        _ods_add_row(funding, [
            row.cycle,
            row.bill_name,
            ("currency", row.due_cents),
            ("currency", row.paid_cents),
            ("currency", row.remaining_cents),
            row.payment_account,
            row.transfer_source,
        ])
    doc.spreadsheet.addElement(funding)

    tx = odf_table.Table(name="Transactions")
    tx_headers = [
        "Date", "Merchant / Description", "Account", "Amount", "Flow",
        "State", "Reconciliation", "Matched Bill",
    ]
    _ods_add_row(tx, tx_headers, header_style="HeaderCell")
    for row in bundle.transactions:
        _ods_add_row(tx, [
            row.date, row.merchant, row.account, ("currency", row.amount_cents),
            "Outflow" if row.amount_cents >= 0 else "Inflow",
            "Pending" if row.pending else "Posted",
            row.reconciliation, row.reconciled_bill,
        ])
    doc.spreadsheet.addElement(tx)

    accounts = odf_table.Table(name="Accounts")
    account_headers = [
        "Account", "Mask", "Type", "Subtype", "Current", "Available", "Bills Checking",
    ]
    _ods_add_row(accounts, account_headers, header_style="HeaderCell")
    for row in bundle.accounts:
        _ods_add_row(accounts, [
            row.name,
            f"••••{row.mask}" if row.mask else None,
            row.account_type, row.account_subtype,
            ("currency", row.current_balance_cents),
            ("currency", row.available_balance_cents),
            "Yes" if row.is_bills_checking else "No",
        ])
    doc.spreadsheet.addElement(accounts)
    doc.save(str(path))


def _pdf_money(cents: int | None) -> str:
    if cents is None:
        return "—"
    value = cents / 100.0
    if value < 0:
        return "-$" + f"{abs(value):,.2f}"
    return "$" + f"{value:,.2f}"


def _pdf_table_style(*, font_size: float = 8) -> TableStyle:
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#10182B")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#C8FF3D")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C8CED8")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
            colors.white, colors.HexColor("#F4F6F8")
        ]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ])


def export_pdf(bundle: ReportBundle, path: Path) -> None:
    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(letter),
        leftMargin=0.4 * inch,
        rightMargin=0.4 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch,
        title=f"Bi-Weekly Bills - {bundle.month_label}",
        author="Bi-Weekly Bills",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        textColor=colors.HexColor("#10182B"),
        fontSize=18,
        leading=21,
        alignment=TA_LEFT,
        spaceAfter=8,
    )
    heading = ParagraphStyle(
        "ReportHeading",
        parent=styles["Heading2"],
        textColor=colors.HexColor("#10182B"),
        spaceBefore=6,
        spaceAfter=6,
    )
    small = ParagraphStyle(
        "ReportSmall",
        parent=styles["BodyText"],
        fontSize=8,
        leading=10,
    )

    story = [
        Paragraph(f"Bi-Weekly Bills — {bundle.month_label}", title_style),
        Paragraph(
            f"Generated {bundle.generated_at.astimezone().strftime('%Y-%m-%d %H:%M:%S %Z')} · "
            f"Bank data: <b>{bundle.bank_label}</b>",
            small,
        ),
        Spacer(1, 8),
    ]

    summary_data = [
        ["Metric", "Month", "1st", "15th"],
        [
            "Scheduled",
            _pdf_money(bundle.month_summary.due_cents),
            _pdf_money(bundle.first_summary.due_cents),
            _pdf_money(bundle.fifteenth_summary.due_cents),
        ],
        [
            "Paid",
            _pdf_money(bundle.month_summary.paid_cents),
            _pdf_money(bundle.first_summary.paid_cents),
            _pdf_money(bundle.fifteenth_summary.paid_cents),
        ],
        [
            "Remaining",
            _pdf_money(bundle.month_summary.remaining_cents),
            _pdf_money(bundle.first_summary.remaining_cents),
            _pdf_money(bundle.fifteenth_summary.remaining_cents),
        ],
        [
            "Bills transfer",
            _pdf_money(bundle.funding.total_transfer_cents),
            _pdf_money(bundle.funding.first_transfer_cents),
            _pdf_money(bundle.funding.fifteenth_transfer_cents),
        ],
    ]
    summary_table = Table(summary_data, colWidths=[1.7 * inch] * 4, repeatRows=1)
    summary_table.setStyle(_pdf_table_style())
    story.extend([summary_table, Spacer(1, 10)])


    story.append(Paragraph("Bills", heading))
    bill_data = [[
        "Cycle", "Bill", "Due", "Paid", "Remaining", "Method", "Status",
        "Payment", "Transfer Source", "Transfer?",
    ]]
    for row in bundle.bills:
        bill_data.append([
            row.cycle, row.name, _pdf_money(row.due_cents),
            _pdf_money(row.paid_cents), _pdf_money(row.remaining_cents),
            row.method or "—", row.status or "—", row.payment_account or "—",
            row.transfer_source or "—", row.transfer_required,
        ])
    bill_table = Table(
        bill_data,
        colWidths=[
            0.48 * inch, 1.35 * inch, 0.72 * inch, 0.72 * inch, 0.78 * inch,
            0.7 * inch, 0.75 * inch, 1.15 * inch, 1.15 * inch, 0.55 * inch,
        ],
        repeatRows=1,
    )
    bill_table.setStyle(_pdf_table_style(font_size=6.8))
    story.append(bill_table)

    story.append(Spacer(1, 10))
    story.append(Paragraph("Bills Checking funding detail", heading))
    if bundle.funding.environment is None:
        story.append(
            Paragraph(
                "Funding assignments are not ready yet.",
                small,
            )
        )
    elif bundle.funding_items:
        funding_data = [[
            "Cycle", "Bill", "Due", "Paid", "Remaining", "Payment Account", "Transfer Source",
        ]]
        for row in bundle.funding_items:
            funding_data.append([
                row.cycle,
                row.bill_name,
                _pdf_money(row.due_cents),
                _pdf_money(row.paid_cents),
                _pdf_money(row.remaining_cents),
                row.payment_account or "—",
                row.transfer_source or "—",
            ])
        funding_table = Table(
            funding_data,
            colWidths=[
                0.55 * inch, 1.75 * inch, 0.85 * inch, 0.85 * inch,
                0.9 * inch, 1.55 * inch, 1.55 * inch,
            ],
            repeatRows=1,
        )
        funding_table.setStyle(_pdf_table_style(font_size=7))
        story.append(funding_table)
    else:
        story.append(Paragraph("No remaining Bills-account funding items.", small))

    story.append(PageBreak())
    story.append(Paragraph("Bank and reconciliation summary", heading))
    story.append(
        Paragraph(
            f"{len(bundle.transactions)} transaction(s) · {bundle.matched_count} matched · "
            f"{bundle.ignored_count} ignored · {bundle.funding_validated_count} funding transfer(s) validated · "
            f"{bundle.unresolved_count} unresolved",
            small,
        )
    )
    story.append(Spacer(1, 6))

    if bundle.accounts:
        account_data = [["Account", "Type", "Current", "Available", "Bills Checking"]]
        for row in bundle.accounts:
            label = row.name + (f" ••••{row.mask}" if row.mask else "")
            account_data.append([
                label,
                " / ".join(
                    value for value in (row.account_type, row.account_subtype) if value
                ) or "—",
                _pdf_money(row.current_balance_cents),
                _pdf_money(row.available_balance_cents),
                "Yes" if row.is_bills_checking else "No",
            ])
        account_table = Table(
            account_data,
            colWidths=[2.4 * inch, 1.5 * inch, 1.2 * inch, 1.2 * inch, 1.0 * inch],
            repeatRows=1,
        )
        account_table.setStyle(_pdf_table_style(font_size=7.5))
        story.extend([account_table, Spacer(1, 10)])

    tx_data = [[
        "Date", "Merchant / Description", "Account", "Amount", "State",
        "Reconciliation", "Matched Bill",
    ]]
    for row in bundle.transactions:
        tx_data.append([
            row.date, row.merchant, row.account, _pdf_money(row.amount_cents),
            "Pending" if row.pending else "Posted", row.reconciliation,
            row.reconciled_bill or "—",
        ])
    tx_table = Table(
        tx_data,
        colWidths=[
            0.78 * inch, 2.05 * inch, 1.7 * inch, 0.8 * inch,
            0.65 * inch, 0.9 * inch, 1.2 * inch,
        ],
        repeatRows=1,
    )
    tx_table.setStyle(_pdf_table_style(font_size=7))
    story.append(tx_table)
    doc.build(story)
